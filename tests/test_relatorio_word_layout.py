"""Testes do layout de fotografias no relatório Word.

Cobrem a escolha de fotos por página (1 a 4) e o tamanho configurável da
foto. Não acessam o banco: o relatório é simulado e as imagens vêm de um
ZIP temporário.
"""

import io
import math
import tempfile
import unittest
import zipfile
from datetime import date
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from relatorio_manutencao.report_generator import gerar_relatorio_word

LARGURA_PX, ALTURA_PX = 300, 200  # 3:2 (paisagem)
FOTOS = 5


def _criar_zip(destino: Path) -> None:
    with zipfile.ZipFile(destino, "w", compression=zipfile.ZIP_DEFLATED) as arquivo:
        for indice in range(FOTOS):
            imagem = Image.new("RGB", (LARGURA_PX, ALTURA_PX), (20 * indice, 90, 160))
            buffer = io.BytesIO()
            imagem.save(buffer, format="JPEG")
            arquivo.writestr(f"FOTO-{indice + 1:02d}.jpg", buffer.getvalue())


def _tem_foto(tabela) -> bool:
    return any(
        celula._tc.findall(".//" + qn("w:drawing"))
        for linha in tabela.rows
        for celula in linha.cells
    )


class LayoutFotosWordTests(unittest.TestCase):
    def setUp(self):
        self.pasta = tempfile.TemporaryDirectory()
        self.origem = Path(self.pasta.name) / "fonte.zip"
        _criar_zip(self.origem)
        self.relatorio = SimpleNamespace(
            caminho_arquivo=str(self.origem),
            data_referencia=date(2026, 10, 27),
            nome_arquivo_original="fonte.zip",
            observacao="",
            periodicidade="MENSAL",
            status="VALIDADO",
        )

    def tearDown(self):
        self.pasta.cleanup()

    def _gerar(self, **opcoes) -> Document:
        caminho = gerar_relatorio_word(
            self.relatorio,
            "Disjuntor",
            "SE Teste",
            "Responsável Teste",
            [],
            caminho_logo=None,
            **opcoes,
        )
        return Document(str(caminho))

    def test_todas_as_fotos_entram_no_documento(self):
        for por_pagina in (1, 2, 3, 4):
            with self.subTest(por_pagina=por_pagina):
                documento = self._gerar(fotos_por_pagina=por_pagina)
                self.assertEqual(len(documento.inline_shapes), FOTOS)
                tabelas = [t for t in documento.tables if _tem_foto(t)]
                self.assertEqual(len(tabelas), math.ceil(FOTOS / por_pagina))

    def test_grade_de_cada_opcao(self):
        esperado = {1: (1, 1), 2: (1, 2), 3: (2, 2), 4: (2, 2)}
        for por_pagina, (linhas, colunas) in esperado.items():
            with self.subTest(por_pagina=por_pagina):
                documento = self._gerar(fotos_por_pagina=por_pagina)
                tabela = next(t for t in documento.tables if _tem_foto(t))
                self.assertEqual(len(tabela.rows), linhas)
                self.assertEqual(len(tabela.columns), colunas)

    def test_foto_unica_ocupa_a_pagina(self):
        # Uma foto por folha deve aproveitar a largura útil (17,4 cm).
        documento = self._gerar(fotos_por_pagina=1)
        self.assertGreater(documento.inline_shapes[0].width.cm, 16)

    def test_tamanho_reduz_a_foto_sem_distorcer(self):
        cheia = self._gerar(fotos_por_pagina=2, tamanho_foto=100).inline_shapes[0]
        reduzida = self._gerar(fotos_por_pagina=2, tamanho_foto=60).inline_shapes[0]
        self.assertLess(reduzida.width.cm, cheia.width.cm)
        for forma in (cheia, reduzida):
            # Regressão: antes a foto era forçada em 4,3 x 9,56 cm (distorcida).
            self.assertAlmostEqual(
                forma.width / forma.height,
                LARGURA_PX / ALTURA_PX,
                places=2,
            )

    def test_valores_fora_da_faixa_sao_ajustados(self):
        documento = self._gerar(fotos_por_pagina=99, tamanho_foto=1)
        self.assertEqual(len(documento.inline_shapes), FOTOS)


if __name__ == "__main__":
    unittest.main()
